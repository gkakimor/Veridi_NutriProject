"""
Escrita padronizada dos workbooks de revisão — openpyxl, sem fórmula e sem macro.

Layout de toda aba: linha 1 = título/legenda, linha 2 = cabeçalho, dados a
partir da linha 3. Cabeçalho verde = coluna editável pela Veridi; cabeçalho
cinza (e célula cinza) = coluna técnica, não editar.
"""

from __future__ import annotations

import math
from dataclasses import dataclass, field
from datetime import date, datetime
from decimal import Decimal
from pathlib import Path

from openpyxl import Workbook
from openpyxl.comments import Comment
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.worksheet.table import Table, TableStyleInfo

FONTE = "Arial"
VERDE = "1E5B3A"
CINZA_CABECALHO = "5F5F5F"
CINZA_CELULA = "EDEDED"
AMARELO = "FFF4CE"
LINHA_CABECALHO = 2
PRIMEIRA_LINHA = 3
FOLGA_VALIDACAO = 300

MOEDA = '"R$" #,##0.00'
MOEDA_PRECISA = '"R$" #,##0.00##'
PERCENTUAL = "0.0##%"
DATA = "DD/MM/YYYY"
TEXTO = "@"


@dataclass
class Coluna:
    nome: str
    descricao: str
    obrigatorio: bool = False
    tecnica: bool = False
    tipo: str = "Texto"
    formato: str = "Texto livre"
    valores: tuple[str, ...] = ()
    lista: str = ""
    pode_vazio: bool = True
    origem: str = ""
    observacao: str = ""
    largura: int = 0
    numero: str = ""
    quebra: bool = False
    validacao_numero: str = ""

    def valores_permitidos(self) -> str:
        if self.valores:
            return ", ".join(v for v in self.valores) + " (ver aba VALORES_PERMITIDOS)"
        if self.lista:
            return f"Chaves existentes ({self.lista}) — lista na própria célula"
        return ""


@dataclass
class Tabela:
    nome_aba: str
    titulo: str
    colunas: list[Coluna]
    linhas: list[dict]
    vazia: str = "Nenhum registro."
    oculta: bool = False
    validar: bool = False
    sufixo_obrigatorio: bool = False
    congelar_coluna: bool = False
    comentarios: bool = False


def cabecalho(coluna: Coluna, tabela: Tabela) -> str:
    return coluna.nome + (" *" if tabela.sufixo_obrigatorio and coluna.obrigatorio else "")


def _valor(valor):
    if valor is None:
        return None
    if isinstance(valor, bool):
        return "SIM" if valor else "NÃO"
    if isinstance(valor, Decimal):
        return float(valor)
    if isinstance(valor, (date, datetime)):
        return valor
    return valor


def _largura(coluna: Coluna, tabela: Tabela) -> int:
    if coluna.largura:
        return coluna.largura
    maior = len(cabecalho(coluna, tabela)) + 3
    for linha in tabela.linhas[:3000]:
        valor = linha.get(coluna.nome)
        if valor is None or valor == "":
            continue
        if isinstance(valor, (date, datetime)):
            tamanho = 11
        elif isinstance(valor, (int, float, Decimal)):
            tamanho = 14
        else:
            tamanho = len(str(valor))
        maior = max(maior, tamanho + 2)
    return max(10, min(maior, 48 if coluna.quebra else 60))


def _escrever(ws, tabela: Tabela, listas: dict[str, str]) -> None:
    fonte = Font(name=FONTE, size=10)
    cinza = PatternFill("solid", fgColor=CINZA_CELULA)
    n_colunas = len(tabela.colunas)
    ultima = get_column_letter(n_colunas)

    larguras = [_largura(c, tabela) for c in tabela.colunas]
    for indice, largura in enumerate(larguras, start=1):
        ws.column_dimensions[get_column_letter(indice)].width = largura

    titulo = ws.cell(row=1, column=1, value=tabela.titulo)
    titulo.font = Font(name=FONTE, size=10, bold=True)
    titulo.fill = PatternFill("solid", fgColor=AMARELO)
    titulo.alignment = Alignment(wrap_text=True, vertical="center")
    if n_colunas > 1:
        ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=n_colunas)
    caracteres_por_linha = max(40, int(sum(larguras) * 1.15))
    linhas_titulo = sum(max(1, math.ceil(len(parte) / caracteres_por_linha)) for parte in tabela.titulo.split("\n"))
    ws.row_dimensions[1].height = min(15 * linhas_titulo + 8, 180)

    for indice, coluna in enumerate(tabela.colunas, start=1):
        celula = ws.cell(row=LINHA_CABECALHO, column=indice, value=cabecalho(coluna, tabela))
        celula.font = Font(name=FONTE, size=10, bold=True, color="FFFFFF")
        celula.fill = PatternFill("solid", fgColor=CINZA_CABECALHO if coluna.tecnica else VERDE)
        celula.alignment = Alignment(wrap_text=True, vertical="center", horizontal="center")
        if tabela.comentarios and coluna.descricao:
            nota = coluna.descricao + (" (Coluna técnica: não editar.)" if coluna.tecnica else "")
            celula.comment = Comment(nota, "Migração Veridi", width=320, height=110)
    ws.row_dimensions[LINHA_CABECALHO].height = 34

    alinhamentos = {
        True: Alignment(wrap_text=True, vertical="top"),
        False: Alignment(wrap_text=False, vertical="top"),
    }
    for numero_linha, linha in enumerate(tabela.linhas, start=PRIMEIRA_LINHA):
        for indice, coluna in enumerate(tabela.colunas, start=1):
            celula = ws.cell(row=numero_linha, column=indice, value=_valor(linha.get(coluna.nome)))
            celula.font = fonte
            celula.alignment = alinhamentos[coluna.quebra]
            if coluna.numero:
                celula.number_format = coluna.numero
            if coluna.tecnica:
                celula.fill = cinza

    ultima_linha = LINHA_CABECALHO + len(tabela.linhas)
    if tabela.linhas:
        tab = Table(displayName=f"tb_{tabela.nome_aba}", ref=f"A{LINHA_CABECALHO}:{ultima}{ultima_linha}")
        tab.tableStyleInfo = TableStyleInfo(name="TableStyleLight1", showRowStripes=True)
        ws.add_table(tab)
    else:
        vazio = ws.cell(row=PRIMEIRA_LINHA, column=1, value=tabela.vazia)
        vazio.font = Font(name=FONTE, size=10, italic=True)
        ws.auto_filter.ref = f"A{LINHA_CABECALHO}:{ultima}{PRIMEIRA_LINHA}"

    ws.freeze_panes = "B3" if tabela.congelar_coluna else "A3"
    ws.page_setup.orientation = "landscape"
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.print_title_rows = f"{LINHA_CABECALHO}:{LINHA_CABECALHO}"

    if not tabela.validar:
        return
    fim = ultima_linha + FOLGA_VALIDACAO
    for indice, coluna in enumerate(tabela.colunas, start=1):
        letra = get_column_letter(indice)
        faixa = f"{letra}{PRIMEIRA_LINHA}:{letra}{fim}"
        validacao = None
        if coluna.valores:
            formula = '"' + ",".join(coluna.valores) + '"'
            if len(formula) > 255:
                raise ValueError(f"Lista longa demais para validação inline: {coluna.nome}")
            validacao = DataValidation(type="list", formula1=formula, allow_blank=coluna.pode_vazio)
            validacao.error = "Use um valor da lista (seta da célula). Significados na aba VALORES_PERMITIDOS."
        elif coluna.lista:
            validacao = DataValidation(type="list", formula1=listas[coluna.lista], allow_blank=coluna.pode_vazio)
            validacao.error = "Use uma chave existente (seta da célula)."
        elif coluna.validacao_numero == "positivo":
            validacao = DataValidation(type="decimal", operator="greaterThan", formula1="0", allow_blank=True)
            validacao.error = "Informe um número maior que zero."
        elif coluna.validacao_numero == "nao_negativo":
            validacao = DataValidation(type="decimal", operator="greaterThanOrEqual", formula1="0", allow_blank=True)
            validacao.error = "Informe um número maior ou igual a zero."
        elif coluna.validacao_numero == "inteiro_positivo":
            validacao = DataValidation(type="whole", operator="greaterThan", formula1="0", allow_blank=True)
            validacao.error = "Informe um número inteiro maior que zero."
        elif coluna.validacao_numero == "fracao":
            validacao = DataValidation(
                type="decimal", operator="between", formula1="0.000001", formula2="1", allow_blank=True
            )
            validacao.error = "Informe o percentual entre 0% e 100% (ex.: 98,5%)."
        elif coluna.tipo == "Data":
            validacao = DataValidation(type="date", operator="greaterThan", formula1="36526", allow_blank=True)
            validacao.error = "Informe uma data no formato dd/mm/aaaa."
        if validacao is not None:
            validacao.errorTitle = "Valor não permitido"
            validacao.showErrorMessage = True
            ws.add_data_validation(validacao)
            validacao.add(faixa)
        if coluna.numero == TEXTO:
            for numero_linha in range(ultima_linha + 1, fim + 1):
                ws.cell(row=numero_linha, column=indice).number_format = TEXTO


def salvar_workbook(caminho: Path, tabelas: list[Tabela], listas: dict[str, list[str]] | None = None, titulo: str = "") -> None:
    wb = Workbook()
    wb.remove(wb.active)
    wb.properties.creator = "Veridi — pacote de migração de cadastros"
    wb.properties.title = titulo or caminho.stem

    referencias: dict[str, str] = {}
    for tabela in tabelas:
        ws = wb.create_sheet(tabela.nome_aba)
        if tabela.oculta:
            ws.sheet_state = "hidden"
        # A aba LISTAS é escrita antes de ser referenciada pelas validações.
        if listas and tabela is tabelas[0]:
            _escrever_listas(wb, listas, referencias)
        _escrever(ws, tabela, referencias)
    if listas:
        wb.move_sheet("LISTAS", offset=len(wb.sheetnames))
    wb.active = 0
    caminho.parent.mkdir(parents=True, exist_ok=True)
    wb.save(caminho)


def _escrever_listas(wb, listas: dict[str, list[str]], referencias: dict[str, str]) -> None:
    ws = wb.create_sheet("LISTAS")
    ws.sheet_state = "hidden"
    for indice, (nome, valores) in enumerate(listas.items(), start=1):
        letra = get_column_letter(indice)
        ws.cell(row=1, column=indice, value=nome).font = Font(name=FONTE, size=10, bold=True)
        for linha, valor in enumerate(valores, start=2):
            ws.cell(row=linha, column=indice, value=valor)
        referencias[nome] = f"LISTAS!${letra}$2:${letra}${max(2, len(valores) + 1)}"
        ws.column_dimensions[letra].width = 40
