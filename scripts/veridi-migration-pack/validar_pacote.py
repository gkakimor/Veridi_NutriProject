#!/usr/bin/env python3
"""
Valida o pacote de revisão da migração — o gerado agora ou o devolvido pela Veridi.

    python scripts/veridi-migration-pack/validar_pacote.py <pasta> [--devolucao]         [--referencia <pasta do pacote anterior>] [--relatorio arquivo.txt]         [--exportar pacote-revisao.json]

ERRO reprova (exit 1); AVISO só informa. Na geração, todo registro nasce com
STATUS_REVISAO = REVISAR e o custo de referência de cada item tem de bater com
os preços dos arquivos 06/07. Com --devolucao, o que depende da decisão da
Veridi (status, custo informado por ela, preço recalculado) vira AVISO.
Com --referencia, compara o pacote com a revisão anterior: nenhuma CHAVE_MIGRACAO
pode sumir, aparecer ou mudar, e nenhuma coluna obrigatória de antes pode ter
sido apagada ou virado opcional. Colunas novas opcionais são permitidas.
Com --exportar, e SÓ quando a validação passa sem erro, grava a leitura
normalizada do pacote (identidade, SHA-256 de cada workbook, e por workbook os
registros indexados pela CHAVE_MIGRACAO). É esse JSON que o pipeline de
migração consome: o Excel entra aqui e em nenhum outro lugar.
Só lê os .xlsx: não conecta em banco nenhum.
"""

from __future__ import annotations

import argparse
import hashlib
import json
import re
import sys
from collections import Counter, defaultdict
from datetime import date, datetime
from decimal import Decimal, InvalidOperation
from pathlib import Path

AQUI = Path(__file__).resolve().parent
sys.path.insert(0, str(AQUI))

from openpyxl import load_workbook  # noqa: E402
from openpyxl.utils import get_column_letter  # noqa: E402

import fontes as F  # noqa: E402
import layout as L  # noqa: E402
import regras as R  # noqa: E402
from planilha import LINHA_CABECALHO, PRIMEIRA_LINHA  # noqa: E402

ABAS = ["DADOS", "DICIONARIO", "VALORES_PERMITIDOS", "PENDENCIAS", "ORIGEM"]
COM_LEGADO = {F.MATERIAS_PRIMAS, F.EMBALAGENS, F.PRODUTOS, F.OFERTAS}
ERRO_EXCEL = re.compile(r"^#(REF!|N/A|VALUE!|DIV/0!|NAME\?|NUM!|NULL!)")
PROIBIDO_06 = re.compile(r"CUSTO_REAL|LAST_REAL|RECEBIMENTO|ACTUAL_UNIT_COST|CUSTO_AQUISICAO", re.I)
PROBLEMAS = {codigo for codigo, *_ in L.PROBLEMAS}
STATUS_INICIAL = "REVISAR"


class Relatorio:
    def __init__(self) -> None:
        self.erros: list[str] = []
        self.avisos: list[str] = []

    def erro(self, arquivo: str, mensagem: str) -> None:
        self.erros.append(f"[{arquivo}] {mensagem}")

    def aviso(self, arquivo: str, mensagem: str) -> None:
        self.avisos.append(f"[{arquivo}] {mensagem}")


def ler(ws) -> tuple[list[str], list[dict]]:
    cabecalho = [str(c.value).removesuffix(" *") if c.value is not None else "" for c in ws[LINHA_CABECALHO]]
    linhas = []
    for numero, valores in enumerate(ws.iter_rows(min_row=PRIMEIRA_LINHA, values_only=True), start=PRIMEIRA_LINHA):
        preenchidos = [v for v in valores if v not in (None, "")]
        if not preenchidos:
            continue
        if len(preenchidos) == 1 and isinstance(valores[0], str) and valores[0].startswith("Nenhum"):
            continue
        linha = dict(zip(cabecalho, valores))
        linha["_linha"] = numero
        linhas.append(linha)
    return cabecalho, linhas


def decimal(valor) -> Decimal | None:
    try:
        return Decimal(str(valor)) if valor not in (None, "") else None
    except InvalidOperation:
        return None


def validar_arquivo(pasta: Path, arquivo: str, rel: Relatorio, devolucao: bool) -> dict | None:
    caminho = pasta / f"{arquivo}.xlsx"
    if not caminho.exists():
        rel.erro(arquivo, "arquivo ausente")
        return None
    try:
        wb = load_workbook(caminho)
    except Exception as falha:  # noqa: BLE001 — qualquer falha de abertura reprova
        rel.erro(arquivo, f"não abre: {falha}")
        return None

    for aba in ABAS:
        if aba not in wb.sheetnames:
            rel.erro(arquivo, f"aba {aba} ausente")
    if arquivo in COM_LEGADO and "LEGADO_NAO_IMPORTADO" not in wb.sheetnames:
        rel.erro(arquivo, "aba LEGADO_NAO_IMPORTADO ausente")
    if wb.sheetnames[0] != "DADOS":
        rel.aviso(arquivo, "a primeira aba não é DADOS")
    if any(a not in wb.sheetnames for a in ABAS):
        return None

    for ws in wb.worksheets:
        for linha in ws.iter_rows():
            for celula in linha:
                if celula.data_type == "f" or (isinstance(celula.value, str) and celula.value.startswith("=")):
                    rel.erro(arquivo, f"fórmula em {ws.title}!{celula.coordinate}")
                elif isinstance(celula.value, str) and ERRO_EXCEL.match(celula.value):
                    rel.erro(arquivo, f"erro do Excel em {ws.title}!{celula.coordinate}: {celula.value}")

    ws = wb["DADOS"]
    specs = L.colunas(arquivo)
    cabecalho, linhas = ler(ws)
    esperado = [s.nome for s in specs]
    if cabecalho[: len(esperado)] != esperado:
        rel.erro(arquivo, f"cabeçalho da aba DADOS difere do esperado: {cabecalho} × {esperado}")
        return None
    for indice, spec in enumerate(specs, start=1):
        bruto = str(ws.cell(row=LINHA_CABECALHO, column=indice).value or "")
        if spec.obrigatorio != bruto.endswith(" *"):
            rel.erro(arquivo, f"marca de obrigatório incoerente em {spec.nome}")

    _, dicionario = ler(wb["DICIONARIO"])
    por_coluna = {d["COLUNA"]: d for d in dicionario}
    for spec in specs:
        entrada = por_coluna.get(spec.nome)
        if not entrada:
            rel.erro(arquivo, f"coluna {spec.nome} sem linha no DICIONARIO")
        elif (entrada["OBRIGATORIO"] == "SIM") != spec.obrigatorio:
            rel.erro(arquivo, f"DICIONARIO diverge da aba DADOS quanto a {spec.nome} ser obrigatório")
    for entrada in dicionario:
        if entrada["OBRIGATORIO"] not in ("SIM", "NÃO") or entrada["PODE_FICAR_VAZIO"] not in ("SIM", "NÃO"):
            rel.erro(arquivo, f"DICIONARIO linha {entrada['_linha']}: OBRIGATORIO/PODE_FICAR_VAZIO devem ser SIM ou NÃO")

    _, valores = ler(wb["VALORES_PERMITIDOS"])
    documentados: dict[str, set] = defaultdict(set)
    for v in valores:
        documentados[v["CAMPO"]].add(v["VALOR"])
    validacoes = ws.data_validations.dataValidation
    for indice, spec in enumerate(specs, start=1):
        if spec.valores:
            faltam = set(spec.valores) - documentados.get(spec.nome, set())
            if faltam:
                rel.erro(arquivo, f"VALORES_PERMITIDOS não documenta {sorted(faltam)} de {spec.nome}")
        if spec.valores or spec.lista:
            celula = f"{get_column_letter(indice)}{PRIMEIRA_LINHA}"
            if not any(celula in v.sqref for v in validacoes):
                rel.erro(arquivo, f"coluna {spec.nome} sem lista suspensa")

    _, pendencias = ler(wb["PENDENCIAS"])
    pend_por_chave: dict[str, list[dict]] = defaultdict(list)
    for p in pendencias:
        if p.get("PROBLEMA") not in PROBLEMAS:
            rel.erro(arquivo, f"PENDENCIAS linha {p['_linha']}: PROBLEMA desconhecido {p.get('PROBLEMA')!r}")
        if p.get("IMPEDE_CARGA") not in ("SIM", "NÃO"):
            rel.erro(arquivo, f"PENDENCIAS linha {p['_linha']}: IMPEDE_CARGA deve ser SIM ou NÃO")
        pend_por_chave[p.get("CHAVE_MIGRACAO")].append(p)

    coluna_status = "STATUS_VALIDACAO" if arquivo == F.MAPA else "STATUS_REVISAO"
    chaves: set[str] = set()
    for linha in linhas:
        numero = linha["_linha"]
        chave = str(linha.get("CHAVE_MIGRACAO") or "")
        if not chave:
            rel.erro(arquivo, f"linha {numero}: CHAVE_MIGRACAO vazia")
            continue
        if not R.PADRAO_CHAVE.match(chave):
            rel.erro(arquivo, f"linha {numero}: CHAVE_MIGRACAO fora do padrão: {chave}")
        if chave in chaves:
            rel.erro(arquivo, f"CHAVE_MIGRACAO repetida: {chave}")
        chaves.add(chave)
        status = linha.get(coluna_status)
        if not devolucao and status != STATUS_INICIAL:
            rel.erro(arquivo, f"{chave}: na geração todo registro nasce {STATUS_INICIAL} (está {status!r})")
        textos = [str(v) for k, v in linha.items() if k != "_linha" and v is not None]
        if any(R.PADRAO_UUID.search(t) for t in textos):
            rel.erro(arquivo, f"linha {numero} ({chave}): UUID técnico no registro")
        texto = " ".join(
            str(v) for k, v in linha.items() if k not in ("_linha", "URL_OU_DOCUMENTO") and isinstance(v, str)
        )
        if R.PADRAO_SINTETICO_CERTO.search(texto):
            rel.erro(arquivo, f"linha {numero} ({chave}): conteúdo reconhecidamente sintético")
        elif R.marcador_sintetico(texto) and arquivo != F.MAPA:
            if not any(p["PROBLEMA"] == "POSSIVEL_DADO_SINTETICO" for p in pend_por_chave.get(chave, [])):
                rel.erro(arquivo, f"linha {numero} ({chave}): possível dado sintético sem pendência")

        for spec in specs:
            valor = linha.get(spec.nome)
            vazio = valor is None or valor == ""
            if spec.obrigatorio and vazio:
                if status == "NAO_IMPORTAR" or pend_por_chave.get(chave):
                    continue
                rel.erro(arquivo, f"linha {numero} ({chave}): obrigatório {spec.nome} vazio sem pendência")
                continue
            if vazio:
                continue
            if spec.valores and str(valor) not in spec.valores:
                rel.erro(arquivo, f"linha {numero} ({chave}): {spec.nome} = {valor!r} fora da lista")
            if spec.tipo == "Data" and not isinstance(valor, (date, datetime)):
                rel.erro(arquivo, f"linha {numero} ({chave}): {spec.nome} não é data")
            if spec.tipo in ("Moeda", "Número", "Percentual", "Inteiro") and not isinstance(valor, (int, float)):
                rel.erro(arquivo, f"linha {numero} ({chave}): {spec.nome} não é número")
            elif spec.validacao_numero == "positivo" and valor <= 0:
                rel.erro(arquivo, f"linha {numero} ({chave}): {spec.nome} deve ser maior que zero")
            elif spec.validacao_numero == "nao_negativo" and valor < 0:
                rel.erro(arquivo, f"linha {numero} ({chave}): {spec.nome} não pode ser negativo")
            elif spec.validacao_numero == "fracao" and not (0 < valor <= 1):
                rel.erro(arquivo, f"linha {numero} ({chave}): {spec.nome} fora de 0–100%")
            elif spec.validacao_numero == "inteiro_positivo" and (int(valor) != valor or valor <= 0):
                rel.erro(arquivo, f"linha {numero} ({chave}): {spec.nome} deve ser inteiro maior que zero")

        bloqueia = any(p["IMPEDE_CARGA"] == "SIM" for p in pend_por_chave.get(chave, []))
        if devolucao and bloqueia and status == "OK":
            rel.aviso(arquivo, f"{chave}: OK com pendência que impedia a carga — conferir se foi resolvida")
        if arquivo == F.CLIENTES and linha.get("CNPJ"):
            digitos = R.so_digitos(str(linha["CNPJ"]))
            if not R.cnpj_valido(digitos) and not any(
                p["PROBLEMA"] == "DOCUMENTO_INVALIDO" for p in pend_por_chave.get(chave, [])
            ):
                (rel.aviso if devolucao else rel.erro)(arquivo, f"{chave}: CNPJ inválido")
        if arquivo in (F.CLIENTES, F.FORNECEDORES) and linha.get("CEP") not in (None, ""):
            if not R.cep_valido(str(linha["CEP"])):
                rel.erro(arquivo, f"linha {numero} ({chave}): CEP {linha['CEP']!r} não tem 8 dígitos "
                                  "(a máscara 00000-000 é aceita; formate a célula como TEXTO para não perder "
                                  "o zero à esquerda)")
        if arquivo == F.PRECOS:
            _validar_preco(arquivo, linha, rel, devolucao)
        if arquivo == F.OFERTAS:
            if linha.get("PRECO") not in (None, "") and not (linha.get("UNIDADE_DO_PRECO") and linha.get("MOEDA")):
                rel.erro(arquivo, f"{chave}: preço sem unidade ou moeda")
            if linha.get("PEDIDO_MINIMO") not in (None, "") and not linha.get("UNIDADE_PEDIDO_MINIMO"):
                rel.erro(arquivo, f"{chave}: pedido mínimo sem unidade")
            inicio, fim = linha.get("VALIDA_A_PARTIR_DE"), linha.get("VALIDADE")
            if isinstance(inicio, (date, datetime)) and isinstance(fim, (date, datetime)) and fim < inicio:
                rel.erro(arquivo, f"{chave}: validade anterior à data de início da oferta")

    if arquivo != F.MAPA:
        for p in pendencias:
            if p.get("CHAVE_MIGRACAO") not in chaves:
                rel.erro(arquivo, f"PENDENCIAS: chave {p.get('CHAVE_MIGRACAO')} não está na aba DADOS")
    if arquivo == F.PRECOS:
        if any(PROIBIDO_06.search(c or "") for c in cabecalho):
            rel.erro(arquivo, "coluna de custo real na planilha de referência de mercado")
        if "NÃO é custo real" not in str(ws.cell(row=1, column=1).value or ""):
            rel.erro(arquivo, "aba DADOS sem o aviso 'NÃO é custo real'")

    coluna_status_spec = next((s for s in specs if s.nome == coluna_status), None)
    return {
        "chaves": chaves,
        "colunas": [s.nome for s in specs],
        "coluna_status": coluna_status,
        "status_permitidos": list(coluna_status_spec.valores) if coluna_status_spec else [],
        "obrigatorias": {s.nome for s in specs if s.obrigatorio},
        "linhas": linhas,
        "status": Counter(l.get(coluna_status) for l in linhas),
        "pendencias": len(pendencias),
        "bloqueantes": sum(1 for p in pendencias if p.get("IMPEDE_CARGA") == "SIM"),
    }


def _validar_preco(arquivo: str, linha: dict, rel: Relatorio, devolucao: bool) -> None:
    chave = linha["CHAVE_MIGRACAO"]
    situacao = linha.get("SITUACAO_PESQUISA")
    preco, quantidade = decimal(linha.get("PRECO_PUBLICADO")), decimal(linha.get("QUANTIDADE_REFERENCIA"))
    unidade = linha.get("UNIDADE_REFERENCIA")
    normalizado = decimal(linha.get("PRECO_NORMALIZADO_POR_UNIDADE_ESTOQUE"))
    if situacao == "ENCONTRADO":
        exigidos = ("PRECO_PUBLICADO", "QUANTIDADE_REFERENCIA", "UNIDADE_REFERENCIA",
                    "PRECO_NORMALIZADO_POR_UNIDADE_ESTOQUE", "URL_OU_DOCUMENTO", "DATA_DA_PESQUISA", "CONFIANCA")
        faltam = [c for c in exigidos if linha.get(c) in (None, "")]
        if faltam:
            rel.erro(arquivo, f"{chave}: ENCONTRADO sem {', '.join(faltam)}")
    if situacao == "NAO_NORMALIZAVEL" and (preco is None or normalizado is not None):
        rel.erro(arquivo, f"{chave}: NAO_NORMALIZAVEL precisa de preço publicado e não pode ter normalizado")
    if situacao in ("SEM_REFERENCIA", "PESQUISA_PENDENTE") and (preco or normalizado or linha.get("CONFIANCA")):
        rel.erro(arquivo, f"{chave}: {situacao} não pode ter preço nem confiança")
    if preco and quantidade and unidade and linha.get("UNIDADE_ESTOQUE"):
        calculado = R.normalizar_preco(preco, quantidade, unidade, linha["UNIDADE_ESTOQUE"])
        if calculado is None and normalizado is not None:
            rel.erro(arquivo, f"{chave}: normalizado sem conversão possível entre {unidade} e {linha['UNIDADE_ESTOQUE']}")
        elif calculado is not None and normalizado is not None:
            tolerancia = max(Decimal("0.0001"), calculado * Decimal("0.000001"))
            if abs(normalizado - calculado) > tolerancia:
                (rel.aviso if devolucao else rel.erro)(arquivo, f"{chave}: normalizado {normalizado} ≠ recalculado {calculado}")


def validar_relacoes(resultados: dict, rel: Relatorio, devolucao: bool) -> None:
    def chaves(arquivo: str) -> set[str]:
        return resultados[arquivo]["chaves"] if resultados.get(arquivo) else set()

    def linhas(arquivo: str) -> list[dict]:
        return (resultados.get(arquivo) or {}).get("linhas", [])

    itens_mp, itens_me = chaves(F.MATERIAS_PRIMAS), chaves(F.EMBALAGENS)
    for chave in itens_mp & itens_me:
        rel.erro("RELACOES", f"item {chave} está em 03 e em 04")
    itens = itens_mp | itens_me
    item_por_chave = {l["CHAVE_MIGRACAO"]: l for l in linhas(F.MATERIAS_PRIMAS) + linhas(F.EMBALAGENS)}
    status_de = {}
    for arquivo in (F.CLIENTES, F.FORNECEDORES, F.MATERIAS_PRIMAS, F.EMBALAGENS, F.PRODUTOS):
        for linha in linhas(arquivo):
            status_de[linha["CHAVE_MIGRACAO"]] = linha.get("STATUS_REVISAO")

    pa: set[str] = set()
    for linha in linhas(F.PRODUTOS):
        chave_pa = str(linha.get("CHAVE_ITEM_PA") or "")
        if chave_pa in pa:
            rel.erro(F.PRODUTOS, f"CHAVE_ITEM_PA repetida: {chave_pa} (regra 1:1)")
        pa.add(chave_pa)
        if chave_pa.removeprefix("PA-LEG-") != str(linha["CHAVE_MIGRACAO"]).removeprefix("PROD-LEG-"):
            rel.erro(F.PRODUTOS, f"{linha['CHAVE_MIGRACAO']}: item de produto acabado {chave_pa} não corresponde (1:1)")
        cliente = linha.get("CHAVE_CLIENTE")
        if cliente and cliente not in chaves(F.CLIENTES):
            rel.erro(F.PRODUTOS, f"{linha['CHAVE_MIGRACAO']}: CHAVE_CLIENTE {cliente} não existe em 01")
        elif cliente and status_de.get(cliente) == "NAO_IMPORTAR" and linha.get("STATUS_REVISAO") != "NAO_IMPORTAR":
            rel.aviso(F.PRODUTOS, f"{linha['CHAVE_MIGRACAO']}: cliente {cliente} marcado NAO_IMPORTAR")

    precos_do_item: dict[str, list[Decimal]] = defaultdict(list)
    cobertos = set()
    for linha in linhas(F.PRECOS):
        cobertos.add(linha.get("CHAVE_ITEM"))
        if linha.get("CHAVE_ITEM") not in itens:
            rel.erro(F.PRECOS, f"{linha['CHAVE_MIGRACAO']}: CHAVE_ITEM {linha.get('CHAVE_ITEM')} não existe em 03/04")
        normalizado = decimal(linha.get("PRECO_NORMALIZADO_POR_UNIDADE_ESTOQUE"))
        if linha.get("SITUACAO_PESQUISA") == "ENCONTRADO" and normalizado is not None \
                and linha.get("STATUS_REVISAO") != "REJEITADA":
            precos_do_item[linha.get("CHAVE_ITEM")].append(normalizado)
    for chave in sorted(itens - cobertos):
        rel.aviso(F.PRECOS, f"item {chave} sem linha de referência de mercado")

    ofertas_do_item: dict[str, list[Decimal]] = defaultdict(list)
    fornecedores_do_item: dict[str, set[str]] = defaultdict(set)
    pares: dict[tuple[str, str], list[dict]] = defaultdict(list)
    for linha in linhas(F.OFERTAS):
        fornecedor, chave_item = linha.get("CHAVE_FORNECEDOR"), linha.get("CHAVE_ITEM")
        if fornecedor not in chaves(F.FORNECEDORES):
            rel.erro(F.OFERTAS, f"{linha['CHAVE_MIGRACAO']}: CHAVE_FORNECEDOR {fornecedor} não existe em 02")
        if chave_item not in itens:
            rel.erro(F.OFERTAS, f"{linha['CHAVE_MIGRACAO']}: CHAVE_ITEM {chave_item} não existe em 03/04")
        for alvo in (fornecedor, chave_item):
            if status_de.get(alvo) == "NAO_IMPORTAR" and linha.get("STATUS_REVISAO") != "NAO_IMPORTAR":
                rel.aviso(F.OFERTAS, f"{linha['CHAVE_MIGRACAO']}: aponta para {alvo}, marcado NAO_IMPORTAR")
        pares[(fornecedor, chave_item)].append(linha)
        fornecedores_do_item[chave_item].add(str(linha.get("FORNECEDOR") or ""))
        item = item_por_chave.get(chave_item)
        preco = decimal(linha.get("PRECO"))
        if item and preco and preco > 0 and linha.get("STATUS_REVISAO") != "NAO_IMPORTAR":
            convertido = R.normalizar_preco(preco, Decimal(1), str(linha.get("UNIDADE_DO_PRECO") or ""), str(item.get("UNIDADE") or ""))
            if convertido is not None:
                ofertas_do_item[chave_item].append(convertido)

    preferidos: Counter = Counter()
    for (fornecedor, chave_item), grupo in pares.items():
        for campo in ("CODIGO_DO_ITEM_NO_FORNECEDOR", "OBSERVACOES_COMERCIAIS", "HOMOLOGACAO",
                      "OBSERVACAO_DA_DECISAO", "PREFERENCIAL", "RELACAO_ATIVA"):
            valores = {str(l.get(campo) or "") for l in grupo}
            if len(valores) > 1:
                rel.erro(F.OFERTAS, f"par {fornecedor} × {chave_item}: {campo} diferente entre as linhas {sorted(valores)}")
        if any(l.get("PREFERENCIAL") == "SIM" for l in grupo):
            preferidos[chave_item] += 1
            if any(l.get("HOMOLOGACAO") != "HOMOLOGADO" for l in grupo):
                rel.erro(F.OFERTAS, f"par {fornecedor} × {chave_item}: preferencial exige HOMOLOGADO")
    for chave_item, quantidade in preferidos.items():
        if quantidade > 1:
            rel.erro(F.OFERTAS, f"item {chave_item}: mais de um fornecedor preferencial")

    for chave, item in item_por_chave.items():
        custo, origem = decimal(item.get("CUSTO_REFERENCIA")), item.get("ORIGEM_CUSTO_REFERENCIA")
        if devolucao and origem == "INFORMADO_PELA_VERIDI":
            continue
        if ofertas_do_item.get(chave):
            esperado, origem_esperada = R.mediana(ofertas_do_item[chave]), "OFERTA_FORNECEDOR_LEGADO"
        elif precos_do_item.get(chave):
            esperado, origem_esperada = R.mediana(precos_do_item[chave]), "PRECO_MERCADO_PUBLICO"
        else:
            esperado, origem_esperada = None, "SEM_REFERENCIA"
        coerente = origem == origem_esperada and (
            (esperado is None and custo is None)
            or (esperado is not None and custo is not None and abs(custo - esperado) <= Decimal("0.000001"))
        )
        if not coerente:
            (rel.aviso if devolucao else rel.erro)(
                "RELACOES", f"{chave}: custo de referência {custo} ({origem}) não bate com 06/07 ({esperado}, {origem_esperada})"
            )
        if not devolucao:
            texto = str(item.get("FORNECEDORES_E_PRECOS_LEGADO") or "")
            for nome in fornecedores_do_item.get(chave, set()):
                if nome and nome not in texto:
                    rel.erro("RELACOES", f"{chave}: oferta de {nome} (07) não aparece em FORNECEDORES_E_PRECOS_LEGADO")

    if resultados.get(F.MAPA):
        mapa = chaves(F.MAPA)
        esperadas = chaves(F.CLIENTES) | chaves(F.FORNECEDORES) | itens | chaves(F.PRODUTOS) | pa
        for chave in sorted(esperadas - mapa):
            rel.erro(F.MAPA, f"registro {chave} fora do mapa de chaves")
        for chave in sorted(mapa - esperadas):
            rel.erro(F.MAPA, f"chave {chave} do mapa não existe nos arquivos de cadastro")


def comparar_com_referencia(pasta: Path, resultados: dict, rel: Relatorio) -> None:
    """Confere o pacote contra a revisão anterior: chaves e colunas obrigatórias.

    O contrato entre revisões é a CHAVE_MIGRACAO. Coluna nova opcional pode
    entrar; chave não pode sumir, aparecer nem mudar, e coluna obrigatória de
    antes não pode desaparecer nem virar opcional.
    """
    for arquivo in F.ARQUIVOS:
        if arquivo not in resultados:
            continue  # comparação parcial (um arquivo só) — não é arquivo ausente
        atual = resultados[arquivo]
        caminho = pasta / f"{arquivo}.xlsx"
        if not caminho.exists():
            rel.aviso(arquivo, f"sem arquivo correspondente na referência ({pasta}) — comparação não feita")
            continue
        if not atual:
            rel.erro(arquivo, "não validado: comparação com a referência não é possível")
            continue
        try:
            wb = load_workbook(caminho, read_only=True)
        except Exception as falha:  # noqa: BLE001 — referência ilegível reprova a comparação
            rel.erro(arquivo, f"referência não abre: {falha}")
            continue
        try:
            ws = wb["DADOS"]
            bruto = [str(c.value) if c.value is not None else "" for c in ws[LINHA_CABECALHO]]
            cabecalho = [c.removesuffix(" *") for c in bruto]
            obrigatorias = {c.removesuffix(" *") for c in bruto if c.endswith(" *")}
            if "CHAVE_MIGRACAO" not in cabecalho:
                rel.erro(arquivo, "referência sem coluna CHAVE_MIGRACAO")
                continue
            indice = cabecalho.index("CHAVE_MIGRACAO")
            antes = {
                str(valores[indice])
                for valores in ws.iter_rows(min_row=PRIMEIRA_LINHA, values_only=True)
                if valores[indice] not in (None, "")
            }
        finally:
            wb.close()

        for coluna in sorted(set(cabecalho) - set(atual["colunas"])):
            if coluna:
                rel.erro(arquivo, f"coluna {coluna} existia na referência e sumiu")
        for coluna in sorted(obrigatorias - atual["obrigatorias"]):
            if coluna in atual["colunas"]:
                rel.erro(arquivo, f"coluna {coluna} era obrigatória na referência e deixou de ser")

        sumiram = sorted(antes - atual["chaves"])
        novas = sorted(atual["chaves"] - antes)
        for chave in sumiram[:20]:
            rel.erro(arquivo, f"CHAVE_MIGRACAO {chave} existia na referência e não está no pacote")
        if len(sumiram) > 20:
            rel.erro(arquivo, f"mais {len(sumiram) - 20} chave(s) da referência ausentes")
        for chave in novas[:20]:
            rel.erro(arquivo, f"CHAVE_MIGRACAO {chave} não existe na referência (chave nova)")
        if len(novas) > 20:
            rel.erro(arquivo, f"mais {len(novas) - 20} chave(s) novas")
        if not sumiram and not novas:
            rel.aviso(arquivo, f"{len(antes)} chaves conferem com a referência (nenhuma nova, nenhuma ausente)")


FORMATO_EXPORTACAO = 1


def _celula(valor):
    """Valor de célula em algo que o JSON aceita. Vazio é sempre None."""
    if valor is None or valor == "":
        return None
    if isinstance(valor, bool):
        return valor
    if isinstance(valor, (int, float)):
        return valor
    if isinstance(valor, datetime):
        return valor.date().isoformat()
    if isinstance(valor, date):
        return valor.isoformat()
    if isinstance(valor, Decimal):
        return str(valor)
    texto = str(valor).strip()
    return texto or None


def _sha256(caminho: Path) -> str:
    digest = hashlib.sha256()
    with caminho.open("rb") as arquivo:
        for bloco in iter(lambda: arquivo.read(1 << 20), b""):
            digest.update(bloco)
    return digest.hexdigest()


def exportar_pacote(pasta: Path, resultados: dict, destino: Path, referencia: Path | None,
                    devolucao: bool) -> dict:
    """Leitura normalizada do pacote, indexada por workbook + CHAVE_MIGRACAO.

    É o contrato entre o Excel e o pipeline da migração: o openpyxl para aqui.
    A identidade do pacote deriva dos SHA-256 dos .xlsx — é com ela que o APPLY
    prova estar aplicando exatamente o pacote que o PLAN aprovou.
    """
    arquivos = []
    for arquivo in F.ARQUIVOS:
        caminho = pasta / f"{arquivo}.xlsx"
        if not caminho.exists():
            continue
        estado = caminho.stat()
        arquivos.append({
            "nome": arquivo,
            "sha256": _sha256(caminho),
            "bytes": estado.st_size,
            "modificadoEm": datetime.fromtimestamp(estado.st_mtime).isoformat(timespec="seconds"),
            "registros": len((resultados.get(arquivo) or {}).get("linhas", [])),
        })

    resumo = '\n'.join(
        f"{a['nome']}:{a['sha256']}" for a in sorted(arquivos, key=lambda a: a["nome"])
    )
    identidade = hashlib.sha256(resumo.encode("utf-8")).hexdigest()

    workbooks = {}
    for arquivo in F.ARQUIVOS:
        r = resultados.get(arquivo)
        if not r:
            continue
        coluna_status = r["coluna_status"]
        registros = []
        for linha in r["linhas"]:
            campos = {
                nome: _celula(linha.get(nome))
                for nome in r["colunas"]
                if nome not in ("CHAVE_MIGRACAO", coluna_status)
            }
            registros.append({
                "chave": str(linha["CHAVE_MIGRACAO"]),
                "status": _celula(linha.get(coluna_status)),
                "campos": campos,
            })
        workbooks[arquivo] = {
            "colunaStatus": coluna_status,
            "statusPermitidos": r["status_permitidos"],
            "colunas": r["colunas"],
            "obrigatorias": sorted(r["obrigatorias"]),
            "contagem": dict(Counter(reg["status"] for reg in registros)),
            "registros": registros,
        }

    pacote = {
        "formato": FORMATO_EXPORTACAO,
        "geradoEm": datetime.now().isoformat(timespec="seconds"),
        "pacote": {
            "caminho": str(pasta.resolve()),
            "identidade": identidade,
            "referencia": str(referencia.resolve()) if referencia else None,
            "arquivos": arquivos,
        },
        "validacao": {"devolucao": devolucao, "erros": 0},
        "workbooks": workbooks,
    }
    destino.parent.mkdir(parents=True, exist_ok=True)
    destino.write_text(json.dumps(pacote, ensure_ascii=False, indent=1) + '\n', encoding="utf-8")
    return pacote


def main() -> int:
    sys.stdout.reconfigure(encoding="utf-8")
    parser = argparse.ArgumentParser(description=__doc__.split("\n\n")[0])
    parser.add_argument("pasta", type=Path)
    parser.add_argument("--devolucao", action="store_true", help="valida arquivos já editados pela Veridi")
    parser.add_argument("--referencia", type=Path,
                        help="pasta do pacote anterior: confere chaves e colunas obrigatórias preservadas")
    parser.add_argument("--exportar", type=Path,
                        help="grava a leitura normalizada do pacote (JSON) quando a validação passa")
    parser.add_argument("--relatorio", type=Path, help="grava o relatório neste arquivo")
    args = parser.parse_args()

    rel = Relatorio()
    resultados = {arquivo: validar_arquivo(args.pasta, arquivo, rel, args.devolucao) for arquivo in F.ARQUIVOS}
    validar_relacoes(resultados, rel, args.devolucao)
    if args.referencia:
        comparar_com_referencia(args.referencia, resultados, rel)

    saida = [f"VALIDAÇÃO DO PACOTE — {args.pasta}", ""]
    if args.referencia:
        saida[1:1] = [f"Referência: {args.referencia}"]
    for arquivo, r in resultados.items():
        if r:
            saida.append(
                f"{arquivo}: {len(r['linhas'])} registros · {r['pendencias']} pendências ({r['bloqueantes']} impedem) · "
                f"{dict(r['status'])}"
            )
        else:
            saida.append(f"{arquivo}: NÃO VALIDADO")
    saida += ["", f"ERROS: {len(rel.erros)}"] + [f"  {e}" for e in rel.erros[:500]]
    saida += ["", f"AVISOS: {len(rel.avisos)}"] + [f"  {a}" for a in rel.avisos[:500]]
    saida += ["", "RESULTADO: " + ("APROVADO" if not rel.erros else "REPROVADO")]

    if args.exportar:
        if rel.erros:
            saida += ["", f"EXPORTAÇÃO: não gerada — {len(rel.erros)} erro(s) na validação."]
        else:
            pacote = exportar_pacote(args.pasta, resultados, args.exportar, args.referencia, args.devolucao)
            saida += [
                "",
                f"EXPORTAÇÃO: {args.exportar}",
                f"  identidade do pacote: {pacote['pacote']['identidade'][:16]}…",
                f"  workbooks exportados: {len(pacote['workbooks'])}",
            ]
    texto = "\n".join(saida)
    print(texto)
    if args.relatorio:
        args.relatorio.write_text(texto + "\n", encoding="utf-8")
    return 1 if rel.erros else 0


if __name__ == "__main__":
    sys.exit(main())
