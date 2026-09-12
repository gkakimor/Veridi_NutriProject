"""
Devolução do 02_FORNECEDORES: os cenários de endereço, chave e status.

    python -m unittest discover -s scripts/veridi-migration-pack

Cada cenário trabalha numa CÓPIA TEMPORÁRIA do workbook gerado — o pacote de
handoff nunca é tocado e nenhum dado de teste sobra nele. Se o pacote não
estiver gerado, os testes são pulados (o corpus real fica fora do Git):

    VERIDI_PACOTE=<pasta>            pacote a conferir (padrão: revisao-02)
    VERIDI_PACOTE_REFERENCIA=<pasta> pacote anterior   (padrão: revisao-01)
"""

from __future__ import annotations

import os
import shutil
import sys
import tempfile
import unittest
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parent))

from openpyxl import load_workbook  # noqa: E402

import fontes as F  # noqa: E402
import validar_pacote as V  # noqa: E402
from planilha import LINHA_CABECALHO, PRIMEIRA_LINHA  # noqa: E402

RAIZ = Path(__file__).resolve().parents[2]
PACOTE = Path(os.environ.get("VERIDI_PACOTE") or RAIZ / "handoff/migracao-producao/revisao-02")
REFERENCIA = Path(os.environ.get("VERIDI_PACOTE_REFERENCIA") or RAIZ / "handoff/migracao-producao/revisao-01")
ARQUIVO = f"{F.FORNECEDORES}.xlsx"
ENDERECO = ("CEP", "LOGRADOURO", "NUMERO", "COMPLEMENTO", "BAIRRO", "CIDADE", "UF")


@unittest.skipUnless((PACOTE / ARQUIVO).exists(), f"pacote não gerado em {PACOTE}")
class Devolucao(unittest.TestCase):
    """Roda o validador em modo devolução sobre uma cópia temporária editada."""

    def validar(self, edicoes: dict[str, object] | None = None, com_referencia: bool = False) -> V.Relatorio:
        """Copia o workbook, aplica as edições na PRIMEIRA linha de dados e valida."""
        with tempfile.TemporaryDirectory(prefix="veridi-devolucao-") as tmp:
            pasta = Path(tmp)
            shutil.copy2(PACOTE / ARQUIVO, pasta / ARQUIVO)
            if edicoes:
                wb = load_workbook(pasta / ARQUIVO)
                ws = wb["DADOS"]
                coluna = {
                    str(c.value).removesuffix(" *"): c.column
                    for c in ws[LINHA_CABECALHO]
                    if c.value is not None
                }
                for nome, valor in edicoes.items():
                    ws.cell(row=PRIMEIRA_LINHA, column=coluna[nome], value=valor)
                wb.save(pasta / ARQUIVO)
            rel = V.Relatorio()
            resultado = V.validar_arquivo(pasta, F.FORNECEDORES, rel, devolucao=True)
            if com_referencia:
                V.comparar_com_referencia(REFERENCIA, {F.FORNECEDORES: resultado}, rel)
            return rel

    def assertSemErro(self, rel: V.Relatorio) -> None:
        self.assertEqual(rel.erros, [], "\n".join(rel.erros))

    def assertErro(self, rel: V.Relatorio, trecho: str) -> None:
        self.assertTrue(
            any(trecho in e for e in rel.erros),
            f"esperava erro contendo {trecho!r}; veio: {rel.erros}",
        )

    # A — endereço vazio é o estado em que o pacote sai: tem de passar.
    def test_a_endereco_vazio_e_valido(self):
        self.assertSemErro(self.validar())

    # B — parcialmente preenchido: só cidade e UF, sem CEP nem logradouro.
    def test_b_endereco_parcial_e_valido(self):
        self.assertSemErro(self.validar({"CIDADE": "Campinas", "UF": "SP"}))

    def test_b2_so_o_cep_tambem_e_valido(self):
        self.assertSemErro(self.validar({"CEP": "13010-000"}))

    # C — endereço completo, CEP com máscara e UF da lista.
    def test_c_endereco_completo_e_valido(self):
        self.assertSemErro(self.validar({
            "CEP": "13010-000",
            "LOGRADOURO": "Avenida Francisco Glicério",
            "NUMERO": "1200",
            "COMPLEMENTO": "Sala 4",
            "BAIRRO": "Centro",
            "CIDADE": "Campinas",
            "UF": "SP",
        }))

    def test_c2_cep_sem_mascara_e_valido(self):
        self.assertSemErro(self.validar({"CEP": "13010000", "CIDADE": "Campinas", "UF": "SP"}))

    # D — CEP inválido reprova.
    def test_d_cep_curto_reprova(self):
        self.assertErro(self.validar({"CEP": "1301"}), "CEP")

    def test_d2_cep_sem_o_zero_a_esquerda_reprova(self):
        # Célula formatada como número perde o zero: 04816-100 vira 4816100.
        self.assertErro(self.validar({"CEP": 4816100}), "CEP")

    # E — UF fora da lista oficial reprova.
    def test_e_uf_invalida_reprova(self):
        self.assertErro(self.validar({"UF": "XX"}), "UF")

    def test_e2_uf_de_pais_estrangeiro_reprova(self):
        self.assertErro(self.validar({"UF": "CA"}), "UF")

    # F — CHAVE_MIGRACAO alterada reprova contra a revisão anterior.
    @unittest.skipUnless((REFERENCIA / ARQUIVO).exists(), f"referência não encontrada em {REFERENCIA}")
    def test_f_chave_alterada_reprova(self):
        rel = self.validar({"CHAVE_MIGRACAO": "FOR-LEG-NOME-TROCADO-NA-DEVOLUCAO"}, com_referencia=True)
        self.assertErro(rel, "existia na referência e não está no pacote")
        self.assertErro(rel, "chave nova")

    @unittest.skipUnless((REFERENCIA / ARQUIVO).exists(), f"referência não encontrada em {REFERENCIA}")
    def test_f2_chave_intacta_passa(self):
        self.assertSemErro(self.validar(com_referencia=True))

    # G — status fora do contrato reprova; os quatro valores aceitos passam.
    def test_g_status_invalido_reprova(self):
        self.assertErro(self.validar({"STATUS_REVISAO": "APROVADO"}), "STATUS_REVISAO")

    def test_g2_status_do_contrato_passam(self):
        for status in ("REVISAR", "OK", "PENDENTE", "NAO_IMPORTAR"):
            with self.subTest(status=status):
                self.assertSemErro(self.validar({"STATUS_REVISAO": status}))

    # Contrato das colunas: endereço opcional, obrigatórios de antes preservados.
    def test_colunas_de_endereco_sao_todas_opcionais(self):
        ws = load_workbook(PACOTE / ARQUIVO, read_only=True)["DADOS"]
        bruto = [str(c.value) if c.value is not None else "" for c in ws[LINHA_CABECALHO]]
        for nome in ENDERECO:
            self.assertIn(nome, [c.removesuffix(" *") for c in bruto])
            self.assertNotIn(f"{nome} *", bruto, f"{nome} não pode ser obrigatório")

    def test_endereco_sai_do_gerador_vazio(self):
        ws = load_workbook(PACOTE / ARQUIVO, read_only=True)["DADOS"]
        cabecalho = [str(c.value).removesuffix(" *") if c.value is not None else "" for c in ws[LINHA_CABECALHO]]
        indices = [cabecalho.index(nome) for nome in ENDERECO]
        for valores in ws.iter_rows(min_row=PRIMEIRA_LINHA, values_only=True):
            if all(v in (None, "") for v in valores):
                continue
            for nome, indice in zip(ENDERECO, indices):
                self.assertIn(valores[indice], (None, ""), f"{nome} veio preenchido pelo gerador")


if __name__ == "__main__":
    unittest.main()
